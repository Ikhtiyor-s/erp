"""
Unit tests for warehouse Excel import/export helpers (apps/api/app/modules/warehouse/excel.py).
These tests exercise the parse/generate layer without a DB connection.

Integration tests (marked with pytest.mark.asyncio) require a running API.
"""

import io
import secrets
from decimal import Decimal

import httpx
import openpyxl
import pytest
import pytest_asyncio

_IMPORT_SUFFIX = secrets.token_hex(4)

from app.modules.warehouse.excel import (
    MAX_FILE_BYTES,
    MAX_ROWS,
    _to_decimal,
    build_export_workbook,
    build_import_template,
    parse_import_file,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_xlsx(rows: list[list]) -> bytes:
    """Build an in-memory .xlsx with the given rows (row 0 = header, rest = data)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = [
    "name", "sku", "barcode", "category", "product_type",
    "unit", "purchase_price", "sale_price",
    "warehouse_name", "opening_qty", "opening_cost", "rack_name",
]


# ---------------------------------------------------------------------------
# parse_import_file
# ---------------------------------------------------------------------------

class TestParseImportFile:
    def test_happy_path_returns_correct_row_count(self):
        xlsx = _make_xlsx([
            HEADER,
            ["Mahsulot A", "SKU-1", "", "Elektronika", "", "dona", "100", "150", "", "", "", ""],
            ["Mahsulot B", "", "", "", "", "", "", "", "", "", "", ""],
        ])
        rows = parse_import_file(xlsx)
        assert len(rows) == 2

    def test_row_numbers_are_one_based_skipping_header(self):
        xlsx = _make_xlsx([
            HEADER,
            ["Alpha", "", "", "", "", "", "", "", "", "", "", ""],
            ["Beta",  "", "", "", "", "", "", "", "", "", "", ""],
        ])
        rows = parse_import_file(xlsx)
        assert rows[0]["_row"] == 2
        assert rows[1]["_row"] == 3

    def test_fields_mapped_correctly(self):
        xlsx = _make_xlsx([
            HEADER,
            ["Widget", "W-001", "1234567890", "Parts", "TypeA",
             "kg", "200", "350", "Main WH", "10", "180", "A-1"],
        ])
        rows = parse_import_file(xlsx)
        r = rows[0]
        assert r["name"] == "Widget"
        assert r["sku"] == "W-001"
        assert r["barcode"] == "1234567890"
        assert r["category"] == "Parts"
        assert r["product_type"] == "TypeA"
        assert r["unit"] == "kg"
        assert r["purchase_price"] == "200"
        assert r["sale_price"] == "350"
        assert r["warehouse_name"] == "Main WH"
        assert r["opening_qty"] == "10"
        assert r["opening_cost"] == "180"
        assert r["rack_name"] == "A-1"

    def test_blank_rows_are_skipped(self):
        xlsx = _make_xlsx([
            HEADER,
            ["Product X", "", "", "", "", "", "", "", "", "", "", ""],
            [None, None, None, None, None, None, None, None, None, None, None, None],
            ["Product Y", "", "", "", "", "", "", "", "", "", "", ""],
        ])
        rows = parse_import_file(xlsx)
        assert len(rows) == 2

    def test_file_too_large_raises_value_error(self):
        oversized = b"x" * (MAX_FILE_BYTES + 1)
        with pytest.raises(ValueError, match="5MB"):
            parse_import_file(oversized)

    def test_returns_empty_list_for_header_only(self):
        xlsx = _make_xlsx([HEADER])
        rows = parse_import_file(xlsx)
        assert rows == []


# ---------------------------------------------------------------------------
# _to_decimal
# ---------------------------------------------------------------------------

class TestToDecimal:
    def test_valid_integer_string(self):
        assert _to_decimal("100", "price") == Decimal("100")

    def test_valid_float_string(self):
        assert _to_decimal("99.99", "price") == Decimal("99.99")

    def test_comma_decimal_separator(self):
        # Comma is treated as decimal separator (European locale), not thousands sep
        assert _to_decimal("1,5", "price") == Decimal("1.5")

    def test_empty_string_returns_zero(self):
        assert _to_decimal("", "price") == Decimal("0")

    def test_non_numeric_raises_value_error(self):
        with pytest.raises(ValueError, match="raqam emas"):
            _to_decimal("arzon", "purchase_price")


# ---------------------------------------------------------------------------
# build_export_workbook
# ---------------------------------------------------------------------------

class TestBuildExportWorkbook:
    def _parse_wb(self, data: bytes) -> openpyxl.Workbook:
        return openpyxl.load_workbook(io.BytesIO(data))

    def test_returns_bytes(self):
        data = build_export_workbook([])
        assert isinstance(data, bytes)
        assert len(data) > 0

    def test_header_row_present(self):
        data = build_export_workbook([])
        wb = self._parse_wb(data)
        ws = wb.active
        assert ws.cell(row=1, column=1).value is not None

    def test_data_row_written_correctly(self):
        rows = [
            {
                "name": "Test Product",
                "sku": "TP-001",
                "barcode": "9876",
                "category_name": "Tools",
                "product_type": "TypeB",
                "unit_name": "dona",
                "purchase_price": Decimal("500"),
                "sale_price": Decimal("700"),
                "warehouse_name": "Asosiy",
                "qty": Decimal("25"),
                "avg_cost": Decimal("500"),
                "rack_name": "B-3",
            }
        ]
        data = build_export_workbook(rows)
        wb = self._parse_wb(data)
        ws = wb.active
        assert ws.cell(row=2, column=1).value == "Test Product"
        assert ws.cell(row=2, column=9).value == "Asosiy"
        assert ws.cell(row=2, column=10).value == 25.0

    def test_twelve_columns_in_header(self):
        data = build_export_workbook([])
        wb = self._parse_wb(data)
        ws = wb.active
        non_empty = sum(
            1 for col in range(1, 13) if ws.cell(row=1, column=col).value is not None
        )
        assert non_empty == 12

    def test_empty_rows_returns_header_only(self):
        data = build_export_workbook([])
        wb = self._parse_wb(data)
        ws = wb.active
        assert ws.max_row == 1


# ---------------------------------------------------------------------------
# build_import_template
# ---------------------------------------------------------------------------

class TestBuildImportTemplate:
    def test_returns_valid_xlsx_bytes(self):
        data = build_import_template()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb is not None

    def test_has_twelve_header_columns(self):
        data = build_import_template()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        non_empty = sum(
            1 for col in range(1, 13) if ws.cell(row=1, column=col).value is not None
        )
        assert non_empty == 12

    def test_first_column_indicates_required(self):
        data = build_import_template()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        assert "*" in str(ws.cell(row=1, column=1).value)


# ---------------------------------------------------------------------------
# HTTP integration tests (require running API + seeded test org)
# ---------------------------------------------------------------------------

XLSX_IMPORT_HEADER = [
    "name", "sku", "barcode", "category", "product_type",
    "unit", "purchase_price", "sale_price",
    "warehouse_name", "opening_qty", "opening_cost", "rack_name",
]


def _make_import_xlsx(rows: list[list]) -> bytes:
    """Build an in-memory .xlsx with IMPORT_HEADER + given data rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(XLSX_IMPORT_HEADER)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _product_row(name: str, sku: str = "", purchase_price: str = "1000") -> list:
    return [name, sku, "", "", "", "dona", purchase_price, "", "", "", "", ""]


@pytest.mark.asyncio
async def test_import_valid_file(client: httpx.AsyncClient):
    """5-row xlsx upload → {created: 5, updated: 0, errors: []}."""
    rows = [
        _product_row(f"ImportTest-T011-{_IMPORT_SUFFIX}-{i}", sku=f"IMP-T011-{_IMPORT_SUFFIX}-{i}")
        for i in range(5)
    ]
    xlsx_bytes = _make_import_xlsx(rows)

    resp = await client.post(
        "/api/v1/warehouse/products/import",
        files={"file": ("products.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["created"] == 5
    assert data["updated"] == 0
    assert data["errors"] == []


@pytest.mark.asyncio
async def test_import_idempotent(client: httpx.AsyncClient):
    """Upload same file twice → second call returns updated: 5."""
    rows = [
        _product_row(f"IdempTest-T011-{_IMPORT_SUFFIX}-{i}", sku=f"IDEMP-T011-{_IMPORT_SUFFIX}-{i}")
        for i in range(5)
    ]
    xlsx_bytes = _make_import_xlsx(rows)

    files = {"file": ("products.xlsx", xlsx_bytes,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r1 = await client.post("/api/v1/warehouse/products/import", files=files)
    assert r1.status_code == 200, r1.text
    assert r1.json()["created"] == 5

    files2 = {"file": ("products.xlsx", xlsx_bytes,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r2 = await client.post("/api/v1/warehouse/products/import", files=files2)
    assert r2.status_code == 200, r2.text
    d2 = r2.json()
    assert d2["updated"] == 5
    assert d2["errors"] == []


@pytest.mark.asyncio
async def test_import_bad_price_row(client: httpx.AsyncClient):
    """One row with non-numeric price → errors[].row set; other rows still processed."""
    rows = [
        _product_row(f"GoodProduct-T011-A-{_IMPORT_SUFFIX}", sku=f"BP-T011-A-{_IMPORT_SUFFIX}"),
        [f"BadPriceProduct-T011-{_IMPORT_SUFFIX}", f"BP-T011-BAD-{_IMPORT_SUFFIX}", "", "", "", "dona", "arzon", "", "", "", "", ""],
        _product_row(f"GoodProduct-T011-B-{_IMPORT_SUFFIX}", sku=f"BP-T011-B-{_IMPORT_SUFFIX}"),
    ]
    xlsx_bytes = _make_import_xlsx(rows)

    resp = await client.post(
        "/api/v1/warehouse/products/import",
        files={"file": ("products.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert len(data["errors"]) >= 1
    error_rows = [e["row"] for e in data["errors"]]
    assert 3 in error_rows, f"Row 3 (bad price) must appear in errors; got {data['errors']}"
    assert data["created"] + data["updated"] >= 2, "Good rows must still be processed"


@pytest.mark.asyncio
async def test_export_returns_xlsx(client: httpx.AsyncClient):
    """GET /warehouse/products/export returns Content-Type xlsx and non-empty body."""
    resp = await client.get("/api/v1/warehouse/products/export")
    assert resp.status_code == 200, resp.text
    ct = resp.headers.get("content-type", "")
    assert "spreadsheetml" in ct, f"Expected xlsx content-type, got: {ct}"
    assert len(resp.content) > 0, "Export file must be non-empty"
